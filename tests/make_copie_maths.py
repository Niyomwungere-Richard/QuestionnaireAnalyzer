"""
Copie de verification MATHEMATIQUES (demande utilisateur).

10 questions, 9 correctes + 1 piege :
  - Q5 (x2 - 9 = 0) : l'etudiant ne donne que x = 3, il manque x = -3
    -> verdict attendu : PARTIEL.

Sortie : documents_tests/copie_D_maths.docx
"""

from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
DOSSIER = BASE_DIR / "documents_tests"
DOSSIER.mkdir(exist_ok=True)

TITRE = "Controle - Mathematiques"
ETUDIANT = "Nom de l'etudiant : Alain N."

QUESTIONS = [
    ("1. Resoudre l'equation suivante : 2x + 6 = 14",
     "Pour trouver x, je soustrais 6 des deux cotes, ce qui donne 2x = 8. "
     "En divisant par 2, j'obtiens x = 4."),
    ("2. Calculer : 5 au carre + 3 au carre",
     "Le carre de 5 vaut 25 et celui de 3 vaut 9. Donc le resultat est 34."),
    ("3. Quelle est la derivee de la fonction f(x) = x au cube ?",
     "La derivee de cette fonction est 3x au carre."),
    ("4. Calculer la moyenne des nombres suivants : 8, 12, 15, 5 et 10.",
     "J'additionne les cinq valeurs pour obtenir 50, puis je divise par 5. "
     "La moyenne est donc 10."),
    ("5. Resoudre l'equation : x au carre - 9 = 0",
     "On ajoute 9 aux deux membres, donc x au carre = 9. Comme 3 au carre "
     "donne 9, la solution est x = 3."),
    ("6. Un triangle possede une base de 10 cm et une hauteur de 6 cm. "
     "Calculer son aire.",
     "L'aire d'un triangle se trouve en multipliant la base par la hauteur "
     "puis en divisant le resultat par deux. J'obtiens donc "
     "10 x 6 / 2 = 30 cm carres."),
    ("7. Developper l'expression suivante : (x + 3)(x + 2)",
     "On distribue chaque terme du premier facteur dans le deuxieme. On "
     "obtient x au carre + 5x + 6."),
    ("8. Si une voiture roule a 60 km/h pendant 2 heures, quelle distance "
     "parcourt-elle ?",
     "La distance se calcule avec vitesse x temps. Donc 60 x 2 = 120 km."),
    ("9. Calculer la probabilite d'obtenir un nombre pair lorsqu'on lance "
     "un de equilibre a six faces.",
     "Les resultats possibles sont 1, 2, 3, 4, 5 et 6. Les nombres pairs "
     "sont 2, 4 et 6. Il y a donc 3 resultats favorables sur 6, soit une "
     "probabilite de 3/6 = 1/2."),
    ("10. Resoudre l'inequation suivante : 3x - 4 > 8",
     "Je commence par ajouter 4 aux deux cotes, ce qui donne 3x > 12. "
     "Ensuite, je divise par 3 et j'obtiens x > 4."),
]


def creer_docx():
    """Cree le Word au format attendu par l'extracteur."""
    import docx

    chemin = DOSSIER / "copie_D_maths.docx"
    doc = docx.Document()
    doc.add_heading(TITRE, level=0)
    doc.add_paragraph(ETUDIANT)
    for question, reponse in QUESTIONS:
        doc.add_heading(question, level=2)
        doc.add_paragraph(f"Reponse : {reponse}")
    doc.save(str(chemin))
    return chemin


if __name__ == "__main__":
    chemin = creer_docx()
    print(f"DOCX cree : {chemin}  ({len(QUESTIONS)} questions)")
    print("Termine !")
