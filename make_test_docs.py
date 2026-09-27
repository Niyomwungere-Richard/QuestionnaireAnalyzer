"""
Creation de documents de test (PDF et Word) pour tester
le module document_handler.

Utilisation :
    python make_test_docs.py

Cree :
    - documents/questionnaire_test.pdf
    - documents/questionnaire_test.docx
"""

from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent
DOSSIER_DOCUMENTS = BASE_DIR / "documents"
DOSSIER_DOCUMENTS.mkdir(exist_ok=True)

# Le texte du questionnaire de test
QUESTIONS_TEST = [
    ("1. Qu'est-ce que TCP ?",
     "TCP est un protocole permettant une communication fiable entre deux machines."),
    ("2. Quelle est la difference entre HTTP et HTTPS ?",
     "HTTPS est la version securisee de HTTP grace au chiffrement SSL/TLS."),
    ("3. Expliquez le modele OSI.",
     "Le modele OSI est un modele en 7 couches qui decrit le fonctionnement des reseaux."),
    ("4. Qu'est-ce qu'une adresse IP ?",
     "Une adresse IP identifie un appareil sur un reseau."),
    ("5. Quel protocole permet de resoudre une adresse de domaine en adresse IP ?",
     ""),  # Pas de reponse = unanswered
]


def creer_pdf():
    """Cree un fichier PDF avec le questionnaire de test."""
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    chemin = DOSSIER_DOCUMENTS / "questionnaire_test.pdf"
    doc = SimpleDocTemplate(str(chemin), pagesize=A4)
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph("Questionnaire de test - Reseaux", styles["Title"]))
    story.append(Spacer(1, 12))

    for question, reponse in QUESTIONS_TEST:
        story.append(Paragraph(question, styles["Heading3"]))
        story.append(Spacer(1, 4))
        if reponse:
            story.append(Paragraph(f"Reponse : {reponse}", styles["Normal"]))
        else:
            story.append(Paragraph("Reponse : ", styles["Normal"]))
        story.append(Spacer(1, 10))

    doc.build(story)
    print(f"PDF cree : {chemin}")
    return chemin


def creer_docx():
    """Cree un fichier Word avec le questionnaire de test."""
    import docx

    chemin = DOSSIER_DOCUMENTS / "questionnaire_test.docx"
    doc = docx.Document()

    doc.add_heading("Questionnaire de test - Reseaux", level=0)

    for question, reponse in QUESTIONS_TEST:
        doc.add_heading(question, level=2)
        if reponse:
            doc.add_paragraph(f"Reponse : {reponse}")
        else:
            doc.add_paragraph("Reponse : ")

    doc.save(str(chemin))
    print(f"DOCX cree : {chemin}")
    return chemin


if __name__ == "__main__":
    print("Creation des documents de test...")
    creer_pdf()
    creer_docx()
    print("Termine !")
