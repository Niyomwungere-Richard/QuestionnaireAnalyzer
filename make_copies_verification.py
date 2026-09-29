"""
Genere 3 copies de verification (ETAPE 17 bis) avec de VRAIS pieges.

Chaque copie teste un comportement precis de l'IA :
  - Copie A (bon eleve)   : reformulations fortes -> doit rester CORRECT
  - Copie B (moyen)       : 1 partiel subtil + 1 incorrect plausible
  - Copie C (faible)      : incorrect + VRAI hors sujet + VRAIE reponse vide

Utilisation :
    venv/Scripts/python.exe make_copies_verification.py

Sorties :
    documents/copie_A_reseaux.docx
    documents/copie_B_systemes.docx
    documents/copie_C_sql.docx
"""

from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent
DOSSIER = BASE_DIR / "documents"
DOSSIER.mkdir(exist_ok=True)

COPIES = {
    "copie_A_reseaux": (
        "Controle - Reseaux TCP/IP",
        "Nom de l'etudiant : Jean-Claude N.",
        [
            # 1. CORRECT fortement reformulee
            ("1. Qu'est-ce qu'une adresse IP et a quoi sert-elle ?",
             "C'est le numero qui permet de reconnaitre chaque appareil "
             "branche sur le reseau, un peu comme une adresse postale : "
             "sans elle, les donnees ne sauraient pas vers quelle machine "
             "se diriger."),
            # 2. CORRECT fortement reformulee
            ("2. Quelle est la difference entre TCP et UDP ?",
             "Avec TCP, les deux ordinateurs se mettent d'accord avant "
             "d'echanger et chaque morceau recu est confirme, donc rien "
             "ne se perd. Avec UDP, on envoie les paquets sans verification, "
             "c'est plus rapide mais certains peuvent ne jamais arriver."),
            # 3. CORRECT
            ("3. Quel est le role du protocole DNS ?",
             "Il convertit un nom de site, facile a retenir pour un humain, "
             "en adresse IP utilisable par les machines."),
            # 4. PARTIEL (bon debut + 1 erreur glissee)
            ("4. Expliquez le fonctionnement general du protocole DHCP.",
             "DHCP attribue automatiquement une adresse IP aux ordinateurs "
             "du reseau au moment ou ils se connectent. C'est le serveur "
             "DHCP qui choisit aussi l'adresse MAC de chaque machine."),
            # 5. CORRECT
            ("5. A quoi sert une passerelle par defaut ?",
             "Quand un paquet doit sortir du reseau local, la machine "
             "l'envoie a la passerelle, qui se charge de le faire suivre "
             "vers les autres reseaux."),
            # 6. CORRECT
            ("6. Qu'est-ce qu'un masque de sous-reseau ?",
             "Il sert a decouper une adresse IP en deux parties : celle "
             "qui designe le reseau et celle qui designe la machine. "
             "Deux machines sont sur le meme sous-reseau si leurs parties "
             "reseau sont identiques."),
        ],
    ),
    "copie_B_systemes": (
        "Controle - Systemes d'exploitation",
        "Nom de l'etudiant : Eric M.",
        [
            # 1. CORRECT
            ("1. Qu'est-ce qu'un systeme d'exploitation ?",
             "C'est le logiciel indispensable qui gere tout l'ordinateur : "
             "il fait le lien entre l'utilisateur, les programmes et le "
             "materiel."),
            # 2. PARTIEL (vrai debut + affirmation trop absolue)
            ("2. Quel est le role du systeme d'exploitation dans la "
             "gestion de la memoire ?",
             "Il donne de la memoire aux programmes qui demarrent et la "
             "reprend quand ils se ferment. Grace a lui, aucun programme "
             "ne peut jamais, dans aucun cas, lire la memoire d'un autre."),
            # 3. CORRECT reformulee
            ("3. Quelle est la difference entre un processus et un "
             "programme ?",
             "Le programme dort sur le disque, c'est juste une suite "
             "d'instructions. Il ne devient un processus qu'au moment ou "
             "on le lance et qu'il s'execute vraiment."),
            # 4. INCORRECT (erreur franche mais plausible)
            ("4. Qu'est-ce qu'un systeme de fichiers ?",
             "C'est le composant qui compresse automatiquement tous les "
             "fichiers du disque pour faire gagner de la place. Sans lui, "
             "le disque serait toujours plein."),
            # 5. CORRECT
            ("5. Expliquez la fonction du noyau d'un systeme "
             "d'exploitation.",
             "C'est le coeur du systeme : il recoit les demandes des "
             "logiciels et pilote directement le processeur, la memoire "
             "et les peripheriques."),
            # 6. CORRECT
            ("6. Qu'est-ce que le multitache ?",
             "Le systeme fait croire que plusieurs programmes tournent en "
             "meme temps en donnant a chacun un petit tour sur le "
             "processeur, l'un apres l'autre."),
        ],
    ),
    "copie_C_sql": (
        "Controle - Bases de donnees SQL",
        "Nom de l'etudiant : Patrick N.",
        [
            # 1. CORRECT simple
            ("1. Qu'est-ce qu'une base de donnees relationnelle ?",
             "Des donnees rangees dans des tables, et ces tables sont "
             "reliees entre elles par des relations."),
            # 2. INCORRECT (inversion classique debutant)
            ("2. Quelle est la difference entre une cle primaire et une "
             "cle etrangere ?",
             "La cle etrangere identifie chaque ligne de facon unique dans "
             "sa table, tandis que la cle primaire sert a se connecter a "
             "une autre table."),
            # 3. PARTIEL (incomplet)
            ("3. A quoi sert la commande SQL SELECT ?",
             "Elle sert a afficher le contenu d'une table."),
            # 4. INCORRECT (confusion des effets)
            ("4. Quelle est la difference entre DELETE et DROP ?",
             "Il n'y a pas de vraie difference : les deux effacent les "
             "donnees, DELETE est juste plus rapide que DROP."),
            # 5. HORS SUJET (repond a cote, sur les index)
            ("5. Qu'est-ce qu'une jointure en SQL ?",
             "Un index permet d'accelerer la recherche dans une table en "
             "evitant de parcourir toutes les lignes une par une."),
            # 6. SANS REPONSE (ligne vraiment vide)
            ("6. A quoi sert la commande GROUP BY ?",
             ""),
        ],
    ),
}


def creer_docx(nom, titre, etudiant, questions):
    """Cree un fichier Word au format attendu par l'extracteur."""
    import docx

    chemin = DOSSIER / f"{nom}.docx"
    doc = docx.Document()
    doc.add_heading(titre, level=0)
    doc.add_paragraph(etudiant)
    for question, reponse in questions:
        doc.add_heading(question, level=2)
        doc.add_paragraph(f"Reponse : {reponse}" if reponse else "Reponse : ")
    doc.save(str(chemin))
    return chemin


if __name__ == "__main__":
    for nom, (titre, etudiant, questions) in COPIES.items():
        chemin = creer_docx(nom, titre, etudiant, questions)
        print(f"DOCX cree : {chemin}  ({len(questions)} questions)")
    print("Termine !")
